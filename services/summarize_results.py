#!/usr/bin/env python3
"""
summarize_results.py

Builds "Summary.xlsx" with two tabs, "Admissions {year}" and "Notes {year}",
summarizing admissions and exam results for the classroom of the given year,
read from the concours_results SQLite database (see init_db.py).

Usage:
    python summarize_results.py 2026
    python summarize_results.py 2026 --output Summary.xlsx

Can also be imported and called as:
    from summarize_results import summarize_results
    summarize_results(2026)
"""

import argparse
from services import db_utils
from config import DB_PATH, OUTPUT_PATH
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"
FONT_SIZE = 10

# Admissions tab: sub-columns under each school, in this exact order.
ADMISSION_ATTR_HEADERS = [
    "Statut",
    "Total points",
    "Moy",
    "Total écrit",
    "Moy écrit",
    "Total oral",
    "Moy oral",
]

STATUS_TO_NA = {"non admissible"}
STATUS_TO_NC = {"non classé-e", "non classee", "non classée", "non classe-e"}
STATUS_TO_RANK = {"admissible", "classé-e", "classee", "classée", "classe-e"}

FORMAT_LABELS = {"written": "écrit", "oral": "oral"}


# --------------------------------------------------------------------------
# Data access
# --------------------------------------------------------------------------

def get_students(conn, classroom_id):
    """All students in the classroom, ordered alphabetically by last name, first name."""
    return conn.execute(
        """
        SELECT cs.id AS classroom_student_id,
               cs.repeating AS repeating,
               s.first_name AS first_name,
               s.last_name AS last_name
        FROM classroom_students cs
        JOIN students s ON s.id = cs.student_id
        WHERE cs.classroom_id = ?
        ORDER BY s.last_name COLLATE NOCASE, s.first_name COLLATE NOCASE
        """,
        (classroom_id,),
    ).fetchall()


def get_admissions(conn, classroom_student_ids):
    if not classroom_student_ids:
        return []
    placeholders = ",".join("?" for _ in classroom_student_ids)
    return conn.execute(
        f"""
        SELECT a.classroom_student_id, a.status, a.rank, a.average,
               a.written_average, a.oral_average, a.total_points,
               a.written_points, a.oral_points,
               sc.id AS school_id, sc.name AS school_name,
               b.id AS bank_id, b.name AS bank_name
        FROM admissions a
        JOIN schools sc ON sc.id = a.school_id
        JOIN banks b ON b.id = sc.bank_id
        WHERE a.classroom_student_id IN ({placeholders})
        """,
        classroom_student_ids,
    ).fetchall()


def get_exam_results(conn, classroom_student_ids):
    if not classroom_student_ids:
        return []
    placeholders = ",".join("?" for _ in classroom_student_ids)
    return conn.execute(
        f"""
        SELECT er.classroom_student_id, er.points,
               e.id AS exam_id, e.name AS exam_name, e.format AS exam_format,
               b.id AS bank_id, b.name AS bank_name
        FROM exam_results er
        JOIN exams e ON e.id = er.exam_id
        JOIN banks b ON b.id = e.bank_id
        WHERE er.classroom_student_id IN ({placeholders})
        """,
        classroom_student_ids,
    ).fetchall()


# --------------------------------------------------------------------------
# Value mapping helpers
# --------------------------------------------------------------------------

def map_status(status, rank):
    if status is None:
        return None
    s = status.strip().lower()
    if s in STATUS_TO_NA:
        return "NA"
    if s in STATUS_TO_NC:
        return "NC"
    if s in STATUS_TO_RANK:
        return rank
    return status


def repeating_label(repeating):
    return "R" if repeating else None


# --------------------------------------------------------------------------
# Styling helpers
# --------------------------------------------------------------------------

def style(cell, bold=False, italic=False, center=False, wrap=False):
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=bold, italic=italic)
    if center or wrap:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)


# --------------------------------------------------------------------------
# Admissions tab
# --------------------------------------------------------------------------

def build_admissions_sheet(wb, year, students, admissions_rows):
    ws = wb.create_sheet(f"Admissions {year}")

    adm_by_key = {}
    schools_by_bank = {}  # bank_name -> {school_id: school_name}
    for r in admissions_rows:
        adm_by_key[(r["classroom_student_id"], r["school_id"])] = r
        schools_by_bank.setdefault(r["bank_name"], {})[r["school_id"]] = r["school_name"]

    bank_names = sorted(schools_by_bank.keys(), key=str.lower)
    columns = []  # (bank_name, school_id, school_name), in display order
    for bank_name in bank_names:
        for school_id, school_name in sorted(schools_by_bank[bank_name].items(), key=lambda kv: kv[1].lower()):
            columns.append((bank_name, school_id, school_name))

    FIRST_DATA_COL = 4  # column D
    N_SUB = len(ADMISSION_ATTR_HEADERS)
    DATA_START_ROW = 12
    n_students = len(students)
    data_end_row = DATA_START_ROW + n_students - 1

    # Fixed labels
    ws["A11"] = "Nom"
    ws["B11"] = "Prénom"
    ws["C11"] = "Redoublant"
    for coord in ("A11", "B11", "C11"):
        style(ws[coord], bold=True)

    for coord, label in (("C3", "Inscrits"), ("C4", "Admissibles"), ("C5", "Pourcentage admissibles"),
                          ("C7", "Moyenne"), ("C8", "Ecart-type"), ("C9", "Min"), ("C10", "Max")):
        ws[coord] = label
        style(ws[coord], italic=True)

    # Per-school blocks
    for idx, (bank_name, school_id, school_name) in enumerate(columns):
        block_start = FIRST_DATA_COL + idx * N_SUB
        block_end = block_start + N_SUB - 1
        start_l = get_column_letter(block_start)
        end_l = get_column_letter(block_end)
        status_l = start_l  # "Statut" is always the first sub-column

        ws.merge_cells(f"{start_l}2:{end_l}2")
        ws[f"{start_l}2"] = school_name
        style(ws[f"{start_l}2"], bold=True, center=True)

        for row_num in (3, 4, 5):
            ws.merge_cells(f"{start_l}{row_num}:{end_l}{row_num}")
        if n_students > 0:
            ws[f"{start_l}3"] = f"=COUNTA({status_l}{DATA_START_ROW}:{status_l}{data_end_row})"
            ws[f"{start_l}4"] = f"=COUNT({status_l}{DATA_START_ROW}:{status_l}{data_end_row})"
            ws[f"{start_l}5"] = f"=IF({start_l}3=0,0,{start_l}4/{start_l}3)"
        for row_num in (3, 4, 5):
            cell = ws[f"{start_l}{row_num}"]
            style(cell, italic=True, center=True)
        ws[f"{start_l}5"].number_format = "0.0%"

        for j, header in enumerate(ADMISSION_ATTR_HEADERS):
            col_l = get_column_letter(block_start + j)
            ws[f"{col_l}6"] = header
            style(ws[f"{col_l}6"], bold=True)

        # Summary stats for the 6 numeric sub-columns (skip Statut)
        for j in range(1, N_SUB):
            col_l = get_column_letter(block_start + j)
            if n_students > 0:
                rng = f"{col_l}{DATA_START_ROW}:{col_l}{data_end_row}"
                ws[f"{col_l}7"] = f'=IFERROR(AVERAGE({rng}),"")'
                ws[f"{col_l}8"] = f'=IFERROR(STDEV({rng}),"")'
                ws[f"{col_l}9"] = f'=IF(COUNT({rng})=0,"",MIN({rng}))'
                ws[f"{col_l}10"] = f'=IF(COUNT({rng})=0,"",MAX({rng}))'
            for row_num in (7, 8, 9, 10):
                style(ws[f"{col_l}{row_num}"], italic=True, center=True)

    # Bank name row (merged across each bank's schools)
    col_cursor = FIRST_DATA_COL
    for bank_name in bank_names:
        n_schools = len(schools_by_bank[bank_name])
        span = n_schools * N_SUB
        start_l = get_column_letter(col_cursor)
        end_l = get_column_letter(col_cursor + span - 1)
        if span > 1:
            ws.merge_cells(f"{start_l}1:{end_l}1")
        ws[f"{start_l}1"] = bank_name
        style(ws[f"{start_l}1"], bold=True, center=True)
        col_cursor += span

    # Student rows
    for i, student in enumerate(students):
        row_num = DATA_START_ROW + i
        ws[f"A{row_num}"] = student["last_name"]
        ws[f"B{row_num}"] = student["first_name"]
        ws[f"C{row_num}"] = repeating_label(student["repeating"])
        for coord in (f"A{row_num}", f"B{row_num}", f"C{row_num}"):
            style(ws[coord])

        for idx, (bank_name, school_id, school_name) in enumerate(columns):
            block_start = FIRST_DATA_COL + idx * N_SUB
            adm = adm_by_key.get((student["classroom_student_id"], school_id))
            values = [None] * N_SUB
            if adm is not None:
                values[0] = map_status(adm["status"], adm["rank"])
                values[1] = adm["total_points"]
                values[2] = adm["average"]
                values[3] = adm["written_points"]
                values[4] = adm["written_average"]
                values[5] = adm["oral_points"]
                values[6] = adm["oral_average"]
            for j, val in enumerate(values):
                col_l = get_column_letter(block_start + j)
                ws[f"{col_l}{row_num}"] = val
                style(ws[f"{col_l}{row_num}"])
                if j > 0:  # skip Statut column
                    ws[f"{col_l}{row_num}"].number_format = "0.00"

    ws.freeze_panes = "D12"
    return ws


# --------------------------------------------------------------------------
# Notes tab
# --------------------------------------------------------------------------

def build_notes_sheet(wb, year, students, exam_result_rows):
    ws = wb.create_sheet(f"Notes {year}")

    results_by_key = {}
    exams_by_bank = {}  # bank_name -> {exam_id: (exam_name, exam_format)}
    for r in exam_result_rows:
        results_by_key[(r["classroom_student_id"], r["exam_id"])] = r["points"]
        exams_by_bank.setdefault(r["bank_name"], {})[r["exam_id"]] = (r["exam_name"], r["exam_format"])

    bank_names = sorted(exams_by_bank.keys(), key=str.lower)
    columns = []  # (bank_name, exam_id, exam_name, exam_format)
    for bank_name in bank_names:
        for exam_id, (exam_name, exam_format) in sorted(exams_by_bank[bank_name].items(), key=lambda kv: kv[1][0].lower()):
            columns.append((bank_name, exam_id, exam_name, exam_format))

    FIRST_DATA_COL = 4  # column D
    DATA_START_ROW = 9
    n_students = len(students)
    data_end_row = DATA_START_ROW + n_students - 1

    ws["A8"] = "Nom"
    ws["B8"] = "Prénom"
    ws["C8"] = "Redoublant"
    for coord in ("A8", "B8", "C8"):
        style(ws[coord], bold=True)

    for coord, label in (("C4", "Moyenne"), ("C5", "Ecart-type"), ("C6", "Min"), ("C7", "Max")):
        ws[coord] = label
        style(ws[coord], italic=True)

    for idx, (bank_name, exam_id, exam_name, exam_format) in enumerate(columns):
        col_l = get_column_letter(FIRST_DATA_COL + idx)

        ws[f"{col_l}2"] = FORMAT_LABELS.get((exam_format or "").strip().lower(), exam_format)
        style(ws[f"{col_l}2"], bold=True)

        ws[f"{col_l}3"] = exam_name
        style(ws[f"{col_l}3"], bold=True, wrap=True)

        if n_students > 0:
            rng = f"{col_l}{DATA_START_ROW}:{col_l}{data_end_row}"
            ws[f"{col_l}4"] = f'=IFERROR(AVERAGE({rng}),"")'
            ws[f"{col_l}5"] = f'=IFERROR(STDEV({rng}),"")'
            ws[f"{col_l}6"] = f'=IF(COUNT({rng})=0,"",MIN({rng}))'
            ws[f"{col_l}7"] = f'=IF(COUNT({rng})=0,"",MAX({rng}))'
        for row_num in (4, 5, 6, 7):
            style(ws[f"{col_l}{row_num}"], italic=True)
            ws[f"{col_l}{row_num}"].number_format = "0.00"

    # Bank name row (merged across each bank's exams)
    col_cursor = FIRST_DATA_COL
    for bank_name in bank_names:
        n_exams = len(exams_by_bank[bank_name])
        start_l = get_column_letter(col_cursor)
        end_l = get_column_letter(col_cursor + n_exams - 1)
        if n_exams > 1:
            ws.merge_cells(f"{start_l}1:{end_l}1")
        ws[f"{start_l}1"] = bank_name
        style(ws[f"{start_l}1"], bold=True)
        col_cursor += n_exams

    for i, student in enumerate(students):
        row_num = DATA_START_ROW + i
        ws[f"A{row_num}"] = student["last_name"]
        ws[f"B{row_num}"] = student["first_name"]
        ws[f"C{row_num}"] = repeating_label(student["repeating"])
        for coord in (f"A{row_num}", f"B{row_num}", f"C{row_num}"):
            style(ws[coord])

        for idx, (bank_name, exam_id, exam_name, exam_format) in enumerate(columns):
            col_l = get_column_letter(FIRST_DATA_COL + idx)
            val = results_by_key.get((student["classroom_student_id"], exam_id))
            ws[f"{col_l}{row_num}"] = val
            style(ws[f"{col_l}{row_num}"])
            ws[f"{col_l}{row_num}"].number_format = "0.00"

    ws.freeze_panes = "D9"
    return ws


# --------------------------------------------------------------------------
# Entry points
# --------------------------------------------------------------------------

def summarize_results(year, output_path=None):
    """
    Build Summary.xlsx for the given classroom year.

    Raises ValueError if no classroom exists for that year.
    Returns the Path to the file written.
    """
    output_path = Path(output_path) if output_path else OUTPUT_PATH

    conn = db_utils.get_connection()
    try:
        with conn:
            classroom_id = db_utils.get_classroom_id_by_year(conn, year)
            students = get_students(conn, classroom_id)
            cs_ids = [s["classroom_student_id"] for s in students]
            admissions_rows = get_admissions(conn, cs_ids)
            exam_result_rows = get_exam_results(conn, cs_ids)
    finally:
        conn.close()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_admissions_sheet(wb, year, students, admissions_rows)
    build_notes_sheet(wb, year, students, exam_result_rows)

    wb.save(output_path)
    print(f"Summarized results into {output_path}")
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Summarize concours results into Summary.xlsx")
    parser.add_argument("year", type=int, help="Classroom year to summarize")
    parser.add_argument("--output", dest="output_path", default=None, help="Output xlsx path (default: Summary.xlsx)")
    args = parser.parse_args()

    try:
        summarize_results(args.year, output_path=args.output_path)
    except ValueError as e:
        raise SystemExit(f"Error: {e}")


if __name__ == "__main__":
    main()
