#!/usr/bin/env python3

from services import excel_utils
from config import STUDENTS_PATH, WRITTEN_STATUSES_PER_BANK_PATH, WRITTEN_EXAM_RESULTS_PATH, WRITTEN_ADMISSIONS_PER_SCHOOL_PATH, ORAL_ADMISSIONS_PER_SCHOOL_PATH, ORAL_STATUSES_PER_BANK_PATH
import openpyxl
from services.anonymize_past_students import anonymize_past_students
from db.init_db import init_db
from services.import_classroom_students import import_classroom_students
from services.import_written_admission_statuses import import_written_admission_statuses
from services.import_written_exam_results import import_written_exam_results
from services.import_written_admission_scores import import_written_admission_scores
from services.import_oral_admission_statuses import import_oral_admission_statuses
from services.import_oral_exam_results_and_admissions import import_oral_exam_results_and_admissions
from services.summarize_results import summarize_results


def validate_classroom_year(path: Path) -> int:
    """Return year given on the tab of the file.

    Raises FileNotFoundError if the file is missing, ValueError if the
    format is invalid (with every issue found listed in the message).
    """
    if not path.exists():
        raise FileNotFoundError(
            f"Missing input file: {path}. Please place the students file there."
        )

    wb = openpyxl.load_workbook(path, data_only=True)

    errors = []

    if len(wb.sheetnames) != 1:
        errors.append(
            f"Expected a single sheet, found {len(wb.sheetnames)}: {wb.sheetnames}"
        )

    sheet_name = wb.sheetnames[0] if wb.sheetnames else None
    year = None
    if sheet_name is not None:
        try:
            year = int(str(sheet_name).strip())
        except ValueError:
            errors.append(
                f"Sheet name '{sheet_name}' is not a valid year (expected e.g. '2025')."
            )

    if errors:
        raise ValueError("Invalid input file format:\n- " + "\n- ".join(errors))

    return year


def validate_written_inputs() -> None:
    excel_utils.discover_files(WRITTEN_STATUSES_PER_BANK_PATH)
    excel_utils.discover_files(WRITTEN_EXAM_RESULTS_PATH)
    excel_utils.discover_files(WRITTEN_ADMISSIONS_PER_SCHOOL_PATH)


def summarize_written_results(year: int) -> None:
    import_classroom_students(year)
    import_written_admission_statuses(year)
    import_written_exam_results(year)
    import_written_admission_scores(year)
    summarize_results(year)


def validate_oral_inputs() -> None:
    excel_utils.discover_files(ORAL_STATUSES_PER_BANK_PATH)
    excel_utils.discover_files(ORAL_ADMISSIONS_PER_SCHOOL_PATH)


def summarize_oral_results(year: int) -> None:
    import_oral_admission_statuses(year)
    import_oral_exam_results_and_admissions(year)
    summarize_results(year)

if __name__ == "__main__":
    try:
        year = validate_classroom_year(STUDENTS_PATH)
        validate_written_inputs()
    except FileNotFoundError as e:
        print(f"Written inputs are invalid: {e}")
        sys.exit(1)

    init_db()
    anonymize_past_students(before=year)
    summarize_written_results(year)

    try:
        validate_oral_inputs()
    except FileNotFoundError:
        print("No oral inputs found. Skipping oral results.")
    else:
        summarize_oral_results(year)
